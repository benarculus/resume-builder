from __future__ import annotations

import importlib.util
import subprocess
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "skills/resume-drafter/scripts/build_docx.py"
FIXTURE = ROOT / "skills/resume-drafter/scripts/fixtures/sample-resume.json"


def test_build_docx_creates_expected_sections(tmp_path: Path) -> None:
    output = tmp_path / "resume.docx"
    subprocess.run(
        [sys.executable, str(SCRIPT), "--input", str(FIXTURE), "--output", str(output)],
        check=True,
    )
    assert output.exists()
    assert output.stat().st_size > 0
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for section in ("Summary", "Experience", "Education", "Skills", "Awards", "Requirements not addressed"):
        assert section in text


def test_job_requirements_parser_accepts_shared_fixture() -> None:
    parser_path = ROOT / "skills/resume-drafter/scripts/parse_job_requirements.py"
    spec = importlib.util.spec_from_file_location("job_requirements_parser", parser_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    artifact = module.parse_job_requirements(
        ROOT / "skills/resume-drafter/scripts/fixtures/sample-job-requirements.json"
    )
    assert artifact["requiredQualifications"][0]["id"] == "RQ-001"


def load_renderer():
    spec = importlib.util.spec_from_file_location("build_docx", SCRIPT)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_build_docx_applies_resume_top_matter_formatting_and_order() -> None:
    renderer = load_renderer()
    document = renderer.build_document(
        {
            "basics": {
                "name": "Jordan Example",
                "email": "jordan@example.com",
                "location": {"city": "Austin", "region": "TX"},
                "securityClearance": "Active Secret",
            },
            "experience": [],
        }
    )

    assert [paragraph.text for paragraph in document.paragraphs[:2]] == [
        "Jordan Example",
        "jordan@example.com | Austin, TX | Active Secret",
    ]
    assert document.paragraphs[0].alignment == 0
    assert document.paragraphs[1].alignment == 0
    assert document.sections[0].left_margin.inches == 0.5
    assert document.sections[0].right_margin.inches == 0.5
    assert document.paragraphs[0].runs[0].font.size.pt == 20


def test_build_docx_omits_unapproved_clearance_from_top_matter() -> None:
    renderer = load_renderer()
    document = renderer.build_document(
        {
            "basics": {
                "name": "Jordan Example",
                "email": "jordan@example.com",
                "location": {"city": "Austin", "region": "TX"},
            },
            "experience": [],
        }
    )

    assert document.paragraphs[1].text == "jordan@example.com | Austin, TX"
    assert "Clearance" not in document.paragraphs[1].text


def test_build_docx_sorts_experience_and_renders_complete_education() -> None:
    renderer = load_renderer()
    document = renderer.build_document(
        {
            "basics": {"name": "Jordan Example"},
            "experience": [
                {"company": "Older Co", "title": "Analyst", "dates": "2018-2020"},
                {"company": "Current Co", "title": "Manager", "dates": "2022-2024"},
            ],
            "education": [
                {
                    "institution": "Example University",
                    "degree": "Bachelor of Science",
                    "major": "Computer Science",
                    "abbreviation": "BSCS",
                    "completionDate": "2024-05",
                }
            ],
        }
    )

    text = [paragraph.text for paragraph in document.paragraphs]
    assert text.index("Manager — Current Co 2022-2024") < text.index(
        "Analyst — Older Co 2018-2020"
    )
    assert (
        "Bachelor of Science in Computer Science (BSCS) — Example University — 2024-05"
        in text
    )


def test_build_docx_renders_start_and_end_dates_when_dates_field_is_absent() -> None:
    renderer = load_renderer()
    document = renderer.build_document(
        {
            "basics": {"name": "Jordan Example"},
            "experience": [
                {
                    "company": "Current Co",
                    "title": "Manager",
                    "startDate": "2022-01",
                    "endDate": "2024-06",
                }
            ],
        }
    )

    assert "Manager — Current Co 2022-01–2024-06" in [
        paragraph.text for paragraph in document.paragraphs
    ]


def test_load_payload_rejects_abbreviation_only_education(tmp_path: Path) -> None:
    renderer = load_renderer()
    payload = tmp_path / "invalid-resume.json"
    payload.write_text(
        '{"basics": {"name": "Jordan Example"}, '
        '"education": [{"institution": "Example University", "abbreviation": "BSCS"}]}',
        encoding="utf-8",
    )

    try:
        renderer.load_payload(payload)
    except ValueError as error:
        assert "full degree or credential" in str(error)
    else:
        raise AssertionError("abbreviation-only education should be rejected")
