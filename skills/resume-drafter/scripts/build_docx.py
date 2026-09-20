#!/usr/bin/env python3
"""Render an approved resume intermediate JSON payload as a portable .docx file."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def load_payload(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as handle:
        payload = json.load(handle)
    if not isinstance(payload, dict):
        raise ValueError("resume input must be a JSON object")
    if not isinstance(payload.get("basics"), dict):
        raise ValueError("resume input must include a basics object")
    experience = payload.get("experience", [])
    if not isinstance(experience, list):
        raise ValueError("resume experience must be a list")
    for index, item in enumerate(experience):
        if not isinstance(item, dict):
            raise ValueError(f"experience entry {index} must be an object")
    education = payload.get("education", [])
    if not isinstance(education, list):
        raise ValueError("resume education must be a list")
    for index, item in enumerate(education):
        if not isinstance(item, dict):
            raise ValueError(f"education entry {index} must be an object")
        if not (item.get("degree") or item.get("studyType")):
            raise ValueError(
                f"education entry {index} must include a full degree or credential"
            )
        if not (item.get("institution") or item.get("organization")):
            raise ValueError(f"education entry {index} must include an institution or provider")
        if not (
            item.get("completionDate")
            or item.get("endDate")
            or item.get("date")
            or item.get("dates")
        ):
            raise ValueError(f"education entry {index} must include a completion date")
    return payload


def add_bullets(document: Document, values: list[str]) -> None:
    for value in values:
        paragraph = document.add_paragraph(str(value), style="List Bullet")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def date_sort_key(value: Any) -> tuple[int, str]:
    text = str(value or "")
    if re.search(r"\bpresent\b", text, re.IGNORECASE):
        return (9999, "12")
    years = re.findall(r"\d{4}", text)
    if not years:
        return (0, "")
    year = int(years[-1])
    if len(years) > 1:
        return (year, "12")
    month_match = re.search(r"\d{4}[-/](\d{1,2})(?!\d)", text)
    month = month_match.group(1) if month_match else None
    return (year, f"{int(month or 0):02d}")


def sorted_entries(entries: list[dict[str, Any]], date_keys: tuple[str, ...]) -> list[dict[str, Any]]:
    return sorted(
        entries,
        key=lambda entry: next(
            (date_sort_key(entry.get(key)) for key in date_keys if entry.get(key)),
            (0, ""),
        ),
        reverse=True,
    )


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    for style_name in ("List Bullet", "List Paragraph"):
        if style_name in document.styles:
            document.styles[style_name].font.name = "Arial"
            document.styles[style_name].font.size = Pt(11)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.bold = True
        run.font.size = Pt(13 if level == 1 else 11)


def contact_values(basics: dict[str, Any]) -> list[str]:
    values = []
    for key in ("email", "location", "securityClearance"):
        value = basics.get(key)
        if isinstance(value, dict) and key == "location":
            value = ", ".join(
                str(part) for part in (value.get("city"), value.get("region")) if part
            )
        if value:
            values.append(str(value))

    # Preserve legacy contact fields after the requested identity/location/clearance order.
    for key in ("phone", "url"):
        if basics.get(key):
            values.append(str(basics[key]))
    return values


def education_details(item: dict[str, Any]) -> str:
    degree = item.get("degree") or item.get("studyType") or ""
    major = item.get("major") or item.get("area") or ""
    abbreviation = item.get("abbreviation") or ""
    qualification = str(degree)
    if major:
        qualification = f"{qualification} in {major}".strip()
    if abbreviation:
        qualification = f"{qualification} ({abbreviation})"
    institution = item.get("institution") or item.get("organization") or ""
    date = (
        item.get("completionDate")
        or item.get("endDate")
        or item.get("date")
        or item.get("dates")
        or ""
    )
    return " — ".join(str(value) for value in (qualification, institution, date) if value)


def role_dates(role: dict[str, Any]) -> str:
    if role.get("dates"):
        return str(role["dates"])
    start = role.get("startDate")
    end = role.get("endDate")
    if start:
        return f"{start}–{end}" if end else str(start)
    if end:
        return str(end)
    return ""


def build_document(payload: dict[str, Any]) -> Document:
    document = Document()
    configure_document(document)
    basics = payload["basics"]
    name = str(basics.get("name", ""))
    heading = document.add_heading(name or "Resume", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in heading.runs:
        run.font.name = "Arial"
        run.font.size = Pt(20)
        run.font.bold = True

    contact = " | ".join(contact_values(basics))
    if contact:
        paragraph = document.add_paragraph(contact)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if payload.get("summary"):
        add_heading(document, "Summary")
        document.add_paragraph(str(payload["summary"]))

    if payload.get("experience"):
        add_heading(document, "Experience")
        roles = payload["experience"]
        ordering = payload.get("experienceOrder", "reverseChronological")
        if ordering not in ("reverseChronological", "approved"):
            raise ValueError(
                "experienceOrder must be 'reverseChronological' or 'approved'"
            )
        ordered_roles = (
            roles
            if ordering == "approved"
            else sorted_entries(roles, ("endDate", "dates", "startDate"))
        )
        for role in ordered_roles:
            title = " — ".join(
                str(value)
                for key in ("title", "company")
                if (value := role.get(key))
            )
            dates = role_dates(role)
            document.add_paragraph(f"{title} {dates}".strip())
            add_bullets(document, [str(item) for item in role.get("bullets", [])])

    if payload.get("education"):
        add_heading(document, "Education")
        items = payload["education"]
        for item in sorted_entries(
            items, ("completionDate", "endDate", "date", "dates")
        ):
            document.add_paragraph(education_details(item))

    if payload.get("skills"):
        add_heading(document, "Skills")
        document.add_paragraph(", ".join(str(item) for item in payload["skills"]))

    if payload.get("awards"):
        add_heading(document, "Awards")
        for item in sorted_entries(payload["awards"], ("date",)):
            details = item.get("details") or item.get("title") or ""
            document.add_paragraph(str(details))

    return document


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    document = build_document(load_payload(args.input))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)


if __name__ == "__main__":
    main()
